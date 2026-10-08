#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert het materiaal voor les 07 "Het digitale klassement" (NovaDepot, ORLO).

  WP2_Het-digitale-klassement.docx   het werkdocument dat de leerling INLEVERT
  nieuw/                             zes nieuwe documenten: van twee nieuwe leveranciers telkens een
                                     leveringsbon, een factuur en een prijslijst
  Nieuwe-documenten.zip              die zes bestanden, zonder mapniveau: dit hang je aan de opdracht
  NovaDepot.zip                      de map NovaDepot zoals ze na les 06 hoort te zijn, voor wie les 06 miste:
                                     ook aan de opdracht hangen

Gebruik:  python3 maak_werkdocumenten.py
Vereist:  python-docx, openpyxl

LET OP bij aanpassen:
 - De zes nieuwe bestanden hebben al een GOEDE naam volgens les 06 (Wat - Van wie). Het eerste woord zegt in
   welke map ze horen (Leveringsbon, Factuur, Prijslijst). Nieuw in deze les is alleen de SLEUTEL vooraan:
   de datum bij leveringsbonnen, het factuurnummer bij facturen. Bij prijslijsten hoeft er niets bij: omdat
   ze allemaal met "Prijslijst - " beginnen, zet Drive ze vanzelf op naam van de leverancier (alfabetisch).
 - De datum op de nieuwe bonnen staat als dag-maand-jaar in cijfers (28-09-2026): de leerling zet de delen
   om naar jaar-maand-dag. De bon van Bakkerij Korstjes (les 06) schrijft de maand voluit; dat is het
   voorbeeld in rij 1 van het werkdocument, en de leraar doet het voor.
 - Op elke factuur staat ook een datum. Die is NIET de sleutel: de leerling kiest het factuurnummer.
 - Zonder sleutel staan de bonnen op naam (Bakkerij Korstjes, Fruitgroothandel Appelbloesem, Zuivelhoeve De
   Melkkan) en dus niet op datum: de nieuwste staat bovenaan. Met de datum vooraan (jaar eerst) klopt het.
 - NovaDepot.zip bevat de vier bestanden van les 06 (rommel/ daar), juist benoemd en in hun map, in
   inhaal/NovaDepot/. Geen mapniveau erboven: "Alles uitpakken" maakt zelf de map NovaDepot (de naam van de
   zip), met Leveringen, Facturen en Prijslijsten erin.
 - Alle bedrijven, namen en adressen zijn fictief. Geen telefoonnummers of e-mailadressen.
"""
import os
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
NIEUW = os.path.join(HERE, "nieuw")
INHAAL = os.path.join(HERE, "inhaal", "NovaDepot")
UIT = os.path.join(HERE, "WP2_Het-digitale-klassement.docx")
ZIP_NIEUW = os.path.join(HERE, "Nieuwe-documenten.zip")
ZIP_INHAAL = os.path.join(HERE, "NovaDepot.zip")

NAVY = RGBColor(0x2A, 0x39, 0x73)
GREY = RGBColor(0x4D, 0x55, 0x73)
GRIJS = RGBColor(0x80, 0x80, 0x80)

# ---------- de gegevens (fictief) ----------
ADRES = {
    "Fruitgroothandel Appelbloesem": "Boomgaardlaan 18, 9820 Merelbeke",
    "Zuivelhoeve De Melkkan": "Weidestraat 40, 9031 Drongen",
}
# wie, levering, poort, rijen (product, aantal, pallets), totaal pallets
BONNEN = [
    ("Fruitgroothandel Appelbloesem", "maandag 28-09-2026, 6.45 uur", "poort 2",
     [("Appels Jonagold, kist van 15 kg", "120", "2"), ("Peren Conference, kist van 15 kg", "60", "1")], "3"),
    ("Zuivelhoeve De Melkkan", "vrijdag 02-10-2026, 7.15 uur", "poort 1",
     [("Volle melk 1 L, tray van 12", "40", "1"), ("Yoghurt naturel 500 g, tray van 12", "30", "1")], "2"),
]
# wie, factuurnummer (de sleutel), datum (niet de sleutel), rijen (product, aantal, prijs, totaal), totaal
FACTUREN = [
    ("Fruitgroothandel Appelbloesem", "2026-0398", "29-09-2026",
     [("Appels Jonagold, kist van 15 kg", "120", "€ 18,50", "€ 2 220,00"),
      ("Peren Conference, kist van 15 kg", "60", "€ 21,00", "€ 1 260,00")], "€ 3 480,00"),
    ("Zuivelhoeve De Melkkan", "2026-0409", "05-10-2026",
     [("Volle melk 1 L, tray van 12", "40", "€ 10,80", "€ 432,00"),
      ("Yoghurt naturel 500 g, tray van 12", "30", "€ 14,40", "€ 432,00")], "€ 864,00"),
]
# wie, rijen (product, inhoud, prijs per stuk)
PRIJSLIJSTEN = [
    ("Fruitgroothandel Appelbloesem",
     [("Appels Jonagold", "kist van 15 kg", 18.50), ("Appels Elstar", "kist van 15 kg", 19.20),
      ("Peren Conference", "kist van 15 kg", 21.00), ("Kiwi's", "doos van 36 stuks", 9.90)]),
    ("Zuivelhoeve De Melkkan",
     [("Volle melk 1 L", "tray van 12", 10.80), ("Halfvolle melk 1 L", "tray van 12", 10.20),
      ("Yoghurt naturel 500 g", "tray van 12", 14.40), ("Jonge kaas, blok van 1 kg", "per stuk", 11.50)]),
]

# Het werkdocument: rij 1 = het voorbeeld uit les 06, al ingevuld; de leerling vult rij 2 en 3 in.
RIJEN_BONNEN = [
    ("Bakkerij Korstjes", "2026-10-13", "2026-10-13 Leveringsbon - Bakkerij Korstjes.docx"),
    ("Fruitgroothandel Appelbloesem", "", ""),
    ("Zuivelhoeve De Melkkan", "", ""),
]
RIJEN_FACTUREN = [
    ("Papierhandel Vellekens", "2026-0417", "2026-0417 Factuur - Papierhandel Vellekens.docx"),
    ("Fruitgroothandel Appelbloesem", "", ""),
    ("Zuivelhoeve De Melkkan", "", ""),
]


# ---------- hulpfuncties (zoals in les 06) ----------
def lettertype(run, naam="Arial"):
    run.font.name = naam
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), naam)


def basis_document(titel):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), "Arial")
    taal = OxmlElement("w:lang")
    taal.set(qn("w:val"), "nl-BE")
    rpr.append(taal)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.1
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
    doc.core_properties.title = titel
    doc.core_properties.author = "Toegepaste Informatica"
    return doc


def tekst(doc, s, vet=False, klein=False, cursief=False, na=6, grootte=11, kleur=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(na)
    r = p.add_run(s)
    lettertype(r)
    r.bold = vet
    r.italic = cursief
    r.font.size = Pt(9.5 if klein else grootte)
    if klein:
        r.font.color.rgb = GREY
    if kleur is not None:
        r.font.color.rgb = kleur
    return p


def kop(doc, s, grootte=13, voor=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(voor)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(s)
    lettertype(r)
    r.bold = True
    r.font.size = Pt(grootte)
    r.font.color.rgb = NAVY
    return p


def _voeg_in_voor(ouder, el, opvolgers):
    """Voeg el in vóór het eerste bestaande element uit opvolgers (OOXML-volgorde), anders achteraan."""
    for tag in opvolgers:
        ref = ouder.find(qn(tag))
        if ref is not None:
            ref.addprevious(el)
            return el
    ouder.append(el)
    return el


def vaste_tabel(t, breedtes_cm, randkleur="8C8C8C"):
    """Vaste kolombreedtes en dunne randen, zodat Google Documenten de tabel toont zoals bedoeld."""
    tbl = t._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = _voeg_in_voor(tblpr, OxmlElement("w:tblW"),
                             ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                              "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    tblw.set(qn("w:w"), str(int(round(sum(breedtes_cm) * 567))))
    tblw.set(qn("w:type"), "dxa")
    randen = OxmlElement("w:tblBorders")
    for kant in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + kant)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), randkleur)
        randen.append(el)
    _voeg_in_voor(tblpr, randen, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"))
    indeling = OxmlElement("w:tblLayout")
    indeling.set(qn("w:type"), "fixed")
    _voeg_in_voor(tblpr, indeling, ("w:tblCellMar", "w:tblLook"))
    for kolom, b in zip(tbl.tblGrid.findall(qn("w:gridCol")), breedtes_cm):
        kolom.set(qn("w:w"), str(int(round(b * 567))))
    for rij in t.rows:
        for cel, b in zip(rij.cells, breedtes_cm):
            cel.width = Cm(b)
    return t


def schaduw(cel, kleur="E4E8F6"):
    tcpr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), kleur)
    tcpr.append(shd)


def cel_tekst(cel, s, vet=False, grootte=10.5, cursief=False):
    cel.text = ""
    p = cel.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(s)
    lettertype(r)
    r.bold = vet
    r.italic = cursief
    r.font.size = Pt(grootte)


def antwoordlijnen(doc, aantal=2):
    for _ in range(aantal):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("_" * 78)
        lettertype(r)
        r.font.color.rgb = GRIJS


def tabel_met_voorbeeld(doc, koppen, breedtes, rijen):
    """Een tabel met kopregel; rij 1 staat al ingevuld (cursief) als voorbeeld."""
    t = doc.add_table(rows=len(rijen) + 1, cols=len(koppen))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    vaste_tabel(t, breedtes)
    for i, k in enumerate(koppen):
        cel_tekst(t.rows[0].cells[i], k, vet=True)
        schaduw(t.rows[0].cells[i])
    for i, rij in enumerate(rijen, start=1):
        for j, w in enumerate(rij):
            if w:
                cel_tekst(t.rows[i].cells[j], w, cursief=(i == 1), grootte=9.5)
        t.rows[i].height = Cm(1.1)
    return t


# ---------- 1. Het werkdocument ----------
def werkdocument():
    doc = basis_document("WP2 Het digitale klassement - werkdocument")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("WERKDOCUMENT — HET DIGITALE KLASSEMENT")
    lettertype(r)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY
    tekst(doc, "Les 07 — Mijn digitale werkplek 2  ·  Toegepaste Informatica  ·  NovaDepot", klein=True, na=8)

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    vaste_tabel(t, [7.0, 4.5, 5.5])
    for i, label in enumerate(["Naam:", "Klas:", "Datum:"]):
        cel_tekst(t.rows[0].cells[i], label + " ", vet=True, grootte=11)

    kop(doc, "Zo werk je", 12, voor=10)
    for s in [
        "1.  Op de lespagina lees je wat je moet doen, stap voor stap.",
        "2.  In Google Drive maak je je klassement. Dat is je echte werk.",
        "3.  In dit werkdocument schrijf je op wat je deed.",
        "4.  Alleen DIT document lever je in.",
    ]:
        tekst(doc, s, na=1)

    # ---- deel 1 ----
    kop(doc, "Deel 1 — Leveringsbonnen: op datum (stap 4)")
    tekst(doc, "Lees de datum bovenaan de bon. Schrijf ze als jaar-maand-dag. Rij 1 is een voorbeeld: "
               "geef die bon in je Drive ook die naam.", klein=True)
    tabel_met_voorbeeld(doc, ["Van wie?", "Datum (jaar-maand-dag)", "Nieuwe naam"], [4.4, 3.6, 9.0], RIJEN_BONNEN)

    # ---- deel 2 ----
    kop(doc, "Deel 2 — Facturen: op nummer (stap 5)")
    tekst(doc, "Lees het factuurnummer bovenaan de factuur. Op de factuur staat ook een datum: die gebruik je "
               "niet. Rij 1 is een voorbeeld.", klein=True)
    tabel_met_voorbeeld(doc, ["Van wie?", "Factuurnummer", "Nieuwe naam"], [4.4, 3.6, 9.0], RIJEN_FACTUREN)

    # ---- deel 3 ----
    kop(doc, "Deel 3 — Drie vragen (stap 6)")
    tekst(doc, "Vraag 1. Je map Prijslijsten staat vanzelf van A tot Z. Hoe heet zo'n klassement? "
               "Waarom moest je daar geen namen veranderen?", vet=True, na=2)
    antwoordlijnen(doc, 2)
    tekst(doc, "Vraag 2. Waarom schrijf je de datum als 2026-10-02 en niet als 02-10-2026?", vet=True, na=2)
    antwoordlijnen(doc, 2)
    tekst(doc, "Vraag 3. Waarom klasseert een bedrijf zijn documenten? Geef één reden.", vet=True, na=2)
    antwoordlijnen(doc, 1)

    # ---- zelfcontrole ----
    kop(doc, "Zelfcontrole — aankruisen vóór je inlevert")
    for s in [
        "In NovaDepot staan geen losse bestanden meer: alles zit in een map.",
        "In Leveringen staan de bonnen met de datum vooraan. De oudste staat bovenaan.",
        "In Facturen staan de facturen met het nummer vooraan. Het kleinste staat bovenaan.",
        "Achter elke naam staat nog .docx of .xlsx.",
        "Deel 1, 2 en 3 zijn ingevuld.",
    ]:
        tekst(doc, "☐  " + s, na=1)

    # ---- extra ----
    kop(doc, "Extra — niet verplicht")
    tekst(doc, "Alleen als je al ingeleverd hebt. Klik in de opdracht op Inleveren ongedaan maken. "
               "Lever daarna opnieuw in.", klein=True)
    tekst(doc, "a) De foto in Leveringen heeft ook een datum: kijk naar het etiket op de foto. "
               "Welke naam geef je de foto?", vet=True, na=2)
    antwoordlijnen(doc, 1)
    tekst(doc, "b) Op 08-10-2026 leverde Sapfabriek Fris. Hoe heet die leveringsbon? "
               "Tussen welke twee bonnen komt hij te staan?", vet=True, na=2)
    antwoordlijnen(doc, 2)

    doc.save(UIT)


# ---------- 2. De zes nieuwe documenten ----------
def bedrijfskop(doc, wat, wie):
    """Bovenaan elk document: WAT in grote letters, daaronder VAN WIE (zoals in les 06)."""
    tekst(doc, wat.upper(), vet=True, grootte=26, kleur=NAVY, na=2)
    tekst(doc, wie, vet=True, grootte=18, na=0)
    tekst(doc, ADRES[wie] + "  ·  fictief bedrijf", klein=True, na=14)


def maak_leveringsbon(wie, levering, poort, rijen, totaal):
    doc = basis_document("Leveringsbon")
    bedrijfskop(doc, "Leveringsbon", wie)
    tekst(doc, f"Aan: NovaDepot, ontvangstzone, {poort}", na=2)
    tekst(doc, f"Levering: {levering}", vet=True, grootte=13, na=10)
    t = doc.add_table(rows=len(rijen) + 2, cols=3)
    t.style = "Table Grid"
    vaste_tabel(t, [8.0, 3.0, 3.0])
    for i, k in enumerate(["Product", "Aantal", "Pallets"]):
        cel_tekst(t.rows[0].cells[i], k, vet=True)
        schaduw(t.rows[0].cells[i])
    for i, rij in enumerate(list(rijen) + [("Totaal", "", totaal)], start=1):
        for j, w in enumerate(rij):
            cel_tekst(t.rows[i].cells[j], w, vet=(i == len(rijen) + 1))
    tekst(doc, "")
    tekst(doc, "Ontvangen door: ............................................", na=2)
    doc.save(os.path.join(NIEUW, f"Leveringsbon - {wie}.docx"))


def maak_factuur(wie, nummer, datum, rijen, totaal):
    doc = basis_document("Factuur")
    bedrijfskop(doc, "Factuur", wie)
    tekst(doc, f"Factuurnummer: {nummer}", vet=True, grootte=13, na=2)
    tekst(doc, f"Datum: {datum}", na=2)
    tekst(doc, "Klant: NovaDepot, Magazijnweg 1, Gent", na=10)
    t = doc.add_table(rows=len(rijen) + 2, cols=4)
    t.style = "Table Grid"
    vaste_tabel(t, [7.0, 2.2, 2.6, 2.6])
    for i, k in enumerate(["Product", "Aantal", "Prijs", "Totaal"]):
        cel_tekst(t.rows[0].cells[i], k, vet=True)
        schaduw(t.rows[0].cells[i])
    for i, rij in enumerate(list(rijen) + [("Totaal zonder btw", "", "", totaal)], start=1):
        for j, w in enumerate(rij):
            cel_tekst(t.rows[i].cells[j], w, vet=(i == len(rijen) + 1))
    tekst(doc, "")
    tekst(doc, "Te betalen binnen 30 dagen.", klein=True)
    doc.save(os.path.join(NIEUW, f"Factuur - {wie}.docx"))


def maak_prijslijst(wie, rijen):
    """Zoals de prijslijst van les 06: PRIJSLIJST, de naam van de leverancier, dan een tabel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Prijslijst"
    ws["A1"] = "PRIJSLIJST"
    ws["A1"].font = Font(name="Arial", size=22, bold=True, color="2A3973")
    ws["A2"] = wie
    ws["A2"].font = Font(name="Arial", size=16, bold=True)
    ws["A3"] = "fictief bedrijf  ·  prijzen voor groothandels, oktober 2026"
    ws["A3"].font = Font(name="Arial", size=9, color="4D5573")
    rand = Border(*[Side(style="thin", color="8C8C8C")] * 4)
    for j, k in enumerate(["Product", "Inhoud", "Prijs per stuk"], start=1):
        c = ws.cell(5, j, k)
        c.font = Font(name="Arial", size=11, bold=True)
        c.fill = PatternFill("solid", fgColor="E4E8F6")
        c.border = rand
    for i, (prod, inh, prijs) in enumerate(rijen, start=6):
        for j, w in enumerate((prod, inh, prijs), start=1):
            c = ws.cell(i, j, w)
            c.font = Font(name="Arial", size=11)
            c.border = rand
            if j == 3:
                c.number_format = '€ #,##0.00'
                c.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    wb.save(os.path.join(NIEUW, f"Prijslijst - {wie}.xlsx"))


def nieuwe_documenten():
    os.makedirs(NIEUW, exist_ok=True)
    for b in BONNEN:
        maak_leveringsbon(*b)
    for f in FACTUREN:
        maak_factuur(*f)
    for p in PRIJSLIJSTEN:
        maak_prijslijst(*p)


# ---------- 3. De zips voor Google Classroom ----------
def zichtbaar(namen):
    return sorted(n for n in namen if not n.startswith("."))   # geen .DS_Store in een zip


def maak_zips():
    # zonder mapniveau: "Alles uitpakken" geeft de map Nieuwe-documenten met de zes bestanden
    namen = zichtbaar(os.listdir(NIEUW))
    with zipfile.ZipFile(ZIP_NIEUW, "w", zipfile.ZIP_DEFLATED) as z:
        for naam in namen:
            z.write(os.path.join(NIEUW, naam), arcname=naam)
    # de drie mappen van NovaDepot, zonder map erboven: "Alles uitpakken" maakt zelf de map NovaDepot
    n = 0
    with zipfile.ZipFile(ZIP_INHAAL, "w", zipfile.ZIP_DEFLATED) as z:
        for map_naam in zichtbaar(os.listdir(INHAAL)):
            map_pad = os.path.join(INHAAL, map_naam)
            if not os.path.isdir(map_pad):
                continue
            for bestand in zichtbaar(os.listdir(map_pad)):
                z.write(os.path.join(map_pad, bestand), arcname=f"{map_naam}/{bestand}")
                n += 1
    return len(namen), n


if __name__ == "__main__":
    werkdocument()
    nieuwe_documenten()
    aantal_nieuw, aantal_inhaal = maak_zips()
    print("Klaar.")
    print("  werkdocument om in te leveren : " + os.path.basename(UIT))
    print("  aan de opdracht te hangen     : " + os.path.basename(ZIP_NIEUW) + f" ({aantal_nieuw} bestanden)")
    print("  voor wie les 06 miste         : " + os.path.basename(ZIP_INHAAL) + f" ({aantal_inhaal} bestanden in 3 mappen)")
